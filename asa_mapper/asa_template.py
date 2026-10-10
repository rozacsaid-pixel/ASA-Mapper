"""The form builder's ASA mapping template (.xlsx): one row per form node and per option; the
"Mapped Asa" column is what the builder imports."""
from __future__ import annotations

import io
import re
import warnings
import zipfile
from typing import Dict, List, Tuple
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from .models import ParsedForm

ALIAS, OPTION, ALIAS_PATH, FIELD_TYPE, MAPPED_ASA = "Form alias", "Field options", "Alias path", "Field type", "Mapped Asa"
REQUIRED = [ALIAS, OPTION, ALIAS_PATH, FIELD_TYPE, MAPPED_ASA]
LAYOUT_TYPES = {"page", "group", "paragraph", "header", "pdf"}
DONE = {"accepted", "changed"}
OPTION_WIDGETS = {"Radio", "MultipleCheckbox"}


class TemplateError(ValueError):
    """The file is not an ASA mapping template, or not one for this form."""


def _load(data: bytes):
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")  # builder files have no default style; harmless
            wb = openpyxl.load_workbook(io.BytesIO(data))
    except Exception as exc:  # openpyxl raises several unrelated types for bad files
        raise TemplateError("The template is not a readable .xlsx.") from exc
    ws = wb.active
    header = [str(c.value or "").strip() for c in ws[1]]
    missing = [c for c in REQUIRED if c not in header]
    if missing:
        raise TemplateError("Not an ASA mapping template; missing columns: " + ", ".join(missing))
    return wb, ws, {name: header.index(name) for name in REQUIRED}


def _row_key(row, col) -> str:
    """'field' for a field row, 'field#option' for an option row, '' for rows to ignore."""
    alias = str(row[col[ALIAS]].value or "").strip()
    option = str(row[col[OPTION]].value or "").strip()
    if alias:
        kind = str(row[col[FIELD_TYPE]].value or "").strip().lower()
        return "" if kind in LAYOUT_TYPES else alias
    if option:
        parts = str(row[col[ALIAS_PATH]].value or "").split("<")
        return f"{parts[-2]}#{option}" if len(parts) >= 2 else ""
    return ""


def template_keys(data: bytes) -> List[str]:
    _, ws, col = _load(data)
    return [k for k in (_row_key(r, col) for r in ws.iter_rows(min_row=2)) if k]


def check_template(data: bytes, form: ParsedForm) -> dict:
    """Compare template rows with the form's fields and options; refuse a template for another form."""
    in_template = set(template_keys(data))
    # Templates list Radio/Checkbox options only, not dropdown or country choices.
    in_form = {f.key for f in form.fields} | {f"{f.key}#{o.key}" for f in form.fields
                                              if f.widget in OPTION_WIDGETS for o in f.options}
    matched = in_form & in_template
    if not matched:
        raise TemplateError("The template does not match this form export (no field keys in common).")
    return {"matched": len(matched), "missing_in_template": sorted(in_form - in_template),
            "not_in_form": sorted(in_template - in_form)}


def final_asas(review: dict, decisions: dict) -> Dict[str, str]:
    """ASA per template key: the reviewer's decision, else the system export's existing ASA.
    PDF-link ASAs are not used as a fallback: they are not what the system has mapped."""
    use_existing = review.get("existing_source") == "asa_export"
    out: Dict[str, str] = {}
    for field in review["fields"]:
        d = decisions.get(field["key"]) or {}
        if d.get("status") in DONE and d.get("asa"):
            out[field["key"]] = d["asa"]
        elif d.get("status") != "none" and use_existing and field["existing_asa"]:
            out[field["key"]] = field["existing_asa"][0]
        for option in field["options"]:
            od = (d.get("options") or {}).get(option["key"]) or {}
            key = f"{field['key']}#{option['key']}"
            if od.get("status") in DONE and od.get("asa"):
                out[key] = od["asa"]
            elif od.get("status") != "none" and use_existing and option["existing_asa"]:
                out[key] = option["existing_asa"][0]
    return out


def fill_template(data: bytes, assignments: Dict[str, str]) -> Tuple[bytes, List[str]]:
    """Write ASAs into the Mapped Asa cells only; every other byte of the file stays as it was
    (re-saving with openpyxl would turn the builder's empty-text cells into missing cells).
    Returns the file and the keys that have no template row."""
    _, ws, col = _load(data)
    letter = get_column_letter(col[MAPPED_ASA] + 1)
    by_row, written = {}, set()
    for row in ws.iter_rows(min_row=2):
        key = _row_key(row, col)
        if key in assignments:
            by_row[row[0].row] = assignments[key]
            written.add(key)

    with zipfile.ZipFile(io.BytesIO(data)) as zin:
        sheet = _active_sheet_path(zin)
        xml = zin.read(sheet).decode("utf-8")
        for number, asa in by_row.items():
            xml = _set_cell(xml, number, letter, asa)
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as zout:
            for info in zin.infolist():
                body = xml.encode("utf-8") if info.filename == sheet else zin.read(info)
                zout.writestr(info, body, compress_type=info.compress_type)
    return out.getvalue(), sorted(set(assignments) - written)


def _active_sheet_path(zin: zipfile.ZipFile) -> str:
    workbook = zin.read("xl/workbook.xml").decode("utf-8")
    rels = zin.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    rid = re.search(r'<sheet\b[^>]*\br:id="([^"]+)"', workbook).group(1)
    target = re.search(r'<Relationship\b[^>]*\bId="%s"[^>]*\bTarget="([^"]+)"' % re.escape(rid), rels)
    if target is None:  # attribute order differs between writers
        target = re.search(r'<Relationship\b[^>]*\bTarget="([^"]+)"[^>]*\bId="%s"' % re.escape(rid), rels)
    path = target.group(1).lstrip("/")
    return path if path.startswith("xl/") else "xl/" + path


def _set_cell(xml: str, number: int, letter: str, value: str) -> str:
    ref = f"{letter}{number}"
    new_cell = lambda style: f'<c r="{ref}" t="inlineStr"{style}><is><t>{escape(value)}</t></is></c>'
    existing = re.search(r'<c r="%s"([^>]*?)(/>|>.*?</c>)' % ref, xml, flags=re.S)
    if existing:
        style = re.search(r'\ss="\d+"', existing.group(1))
        return xml[:existing.start()] + new_cell(style.group(0) if style else "") + xml[existing.end():]
    row = re.search(r'<row r="%d"([^>]*?)(/>|>(.*?)</row>)' % number, xml, flags=re.S)
    if row is None:
        raise TemplateError(f"Row {number} not found in the template sheet.")
    if row.group(2) == "/>":
        return xml[:row.start()] + f'<row r="{number}"{row.group(1)}>{new_cell("")}</row>' + xml[row.end():]
    cells = row.group(3)
    target = column_index_from_string(letter)
    insert_at = len(cells)
    for m in re.finditer(r'<c r="([A-Z]+)\d+"', cells):
        if column_index_from_string(m.group(1)) > target:
            insert_at = m.start()
            break
    cells = cells[:insert_at] + new_cell("") + cells[insert_at:]
    return xml[:row.start()] + f'<row r="{number}"{row.group(1)}>{cells}</row>' + xml[row.end():]
