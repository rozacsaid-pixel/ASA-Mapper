"""Load the ASA Library ("ASA List" tab) from xlsx bytes or Google Sheets."""
from __future__ import annotations

import io
import json
from pathlib import Path

import openpyxl

from .models import AsaEntry, Library

DEFAULT_TAB = "ASA List"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
REQUIRED_COLUMNS = ["Section", "Field meaning summary", "Field name", "Field option", "Type",
                    "Available under field name"]
OPTIONAL_COLUMNS = ["Field group", "Field meaning full definition", "Keywords", "Input type in the form"]
KINDS = {"field name": "name", "field option": "option"}


class LibraryError(ValueError):
    """The library file cannot be used."""


def _s(value) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def load_library_xlsx(data: bytes, tab: str = DEFAULT_TAB) -> Library:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types for bad files
        raise LibraryError("The library file is not a readable .xlsx.") from exc
    if tab not in wb.sheetnames:
        raise LibraryError(f"Tab '{tab}' not found. Tabs: {', '.join(wb.sheetnames)}")
    rows = list(wb[tab].iter_rows(values_only=True))

    header_idx = None
    for i, row in enumerate(rows[:10]):
        cells = [_s(c) for c in row or ()]
        if "Field name" in cells and "Type" in cells:
            header_idx = i
            break
    if header_idx is None:
        raise LibraryError("Header row with 'Field name' and 'Type' not found in the first 10 rows.")
    header = [_s(c) for c in rows[header_idx]]
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise LibraryError("Missing columns: " + ", ".join(missing))
    col = {name: header.index(name) for name in REQUIRED_COLUMNS + OPTIONAL_COLUMNS if name in header}

    def get(row, name):
        i = col.get(name)
        return _s(row[i]) if i is not None and i < len(row) else ""

    entries, issues, seen = [], [], set()
    for number, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not row or not any(row):
            continue
        kind = KINDS.get(get(row, "Type").lower())
        asa = get(row, "Field name" if kind == "name" else "Field option") if kind else ""
        if not kind or not asa.startswith("asa_"):
            issues.append(f"Row {number}: skipped (Type '{get(row, 'Type')}', alias '{asa}')")
            continue
        if asa in seen:
            issues.append(f"Row {number}: duplicate {asa} skipped")
            continue
        seen.add(asa)
        entries.append(AsaEntry(
            asa=asa, kind=kind, section=get(row, "Section"), group=get(row, "Field group"),
            summary=get(row, "Field meaning summary"), definition=get(row, "Field meaning full definition"),
            keywords=get(row, "Keywords"),
            parent=get(row, "Available under field name") if kind == "option" else "",
            input_type=get(row, "Input type in the form"),
        ))

    names = {e.asa for e in entries if e.kind == "name"}
    for e in entries:
        if e.kind == "option" and not e.parent:
            issues.append(f"{e.asa}: option has no parent field name")
        elif e.kind == "option" and e.parent not in names:
            issues.append(f"{e.asa}: parent {e.parent} not in library")
    return Library(entries=entries, issues=issues)


def service_account_info(value: str) -> dict:
    value = value.strip()
    if value.startswith("{"):
        return json.loads(value)
    return json.loads(Path(value).read_text())


def fetch_google_sheet_xlsx(sheet_id: str, service_account: str) -> bytes:
    """Download the whole sheet as .xlsx through the Drive API (sheet must be shared with the account)."""
    from google.auth.transport.requests import AuthorizedSession
    from google.oauth2 import service_account as sa

    creds = sa.Credentials.from_service_account_info(
        service_account_info(service_account),
        scopes=["https://www.googleapis.com/auth/drive.readonly"])
    response = AuthorizedSession(creds).get(
        f"https://www.googleapis.com/drive/v3/files/{sheet_id}/export",
        params={"mimeType": XLSX_MIME}, timeout=60)
    response.raise_for_status()
    return response.content
