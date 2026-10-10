import io

import openpyxl

from asa_mapper.export_parser import parse_export
from asa_mapper.exporter import SHEET_HEADERS, review_sheet
from asa_mapper.matcher import Matcher
from asa_mapper.service import build_review


def make_review(export_dir, library):
    form = parse_export(export_dir)
    return form, build_review(form, Matcher(library), library)


def test_build_review(export_dir, library):
    _, review = make_review(export_dir, library)
    sug = review["suggestions"]
    assert "sig" not in sug and "split1" not in sug
    assert sug["phone1"]["field"][0]["asa"] == "asa_phone_primarycontact_contactinfo"
    assert sug["email1"]["field"][0]["asa"] == "asa_email_primarycontact_contactinfo"
    assert sug["bank"]["parent"] == "asa_bhc_specialentities"
    assert sug["bank"]["options"]["Yes"][0]["asa"] == "asa_yes_bhc_specialentities"
    assert sug["phone1"]["options"] == {}
    assert "not in library: asa_commitment_amount" in sug["commitment"]["flags"]
    assert len(review["fields"]) == 9


def test_review_sheet_rows(export_dir, library):
    form, review = make_review(export_dir, library)
    decisions = {"bank": {"status": "accepted", "asa": "asa_bhc_specialentities", "reviewer": "me@x",
                          "at": "2026-10-10T10:00:00+00:00",
                          "options": {"Yes": {"status": "accepted", "asa": "asa_yes_bhc_specialentities"}}}}
    ws = openpyxl.load_workbook(io.BytesIO(review_sheet(form.name, review, decisions))).active
    assert [c.value for c in ws[1]] == SHEET_HEADERS
    rows = [dict(zip(SHEET_HEADERS, [c.value or "" for c in r])) for r in ws.iter_rows(min_row=2)]
    bank = next(r for r in rows if r["Field key"] == "bank" and not r["Option key"])
    assert (bank["Accepted ASA"], bank["Status"], bank["Reviewer"]) == ("asa_bhc_specialentities", "accepted", "me@x")
    yes = next(r for r in rows if r["Field key"] == "bank" and r["Option key"] == "Yes")
    assert yes["Suggested ASA"] == yes["Accepted ASA"] == "asa_yes_bhc_specialentities"
    sig = next(r for r in rows if r["Field key"] == "sig")
    assert sig["Status"] == "skipped" and "Signature" in sig["Flags"]
    assert any(r["Field key"] == "commitment" and "not in library" in r["Flags"] for r in rows)
    assert next(r for r in rows if r["Field key"] == "phone1")["Status"] == "pending"


def test_review_sheet_keeps_decision_on_skipped_field(export_dir, library):
    form, review = make_review(export_dir, library)
    decisions = {"sig": {"status": "changed", "asa": "asa_fullname_investorname_generalinfo"}}
    ws = openpyxl.load_workbook(io.BytesIO(review_sheet(form.name, review, decisions))).active
    rows = [dict(zip(SHEET_HEADERS, [c.value or "" for c in r])) for r in ws.iter_rows(min_row=2)]
    sig = next(r for r in rows if r["Field key"] == "sig")
    assert (sig["Accepted ASA"], sig["Status"]) == ("asa_fullname_investorname_generalinfo", "changed")
    assert "Signature" in sig["Flags"]
