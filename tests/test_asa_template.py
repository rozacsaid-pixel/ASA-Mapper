import io

import openpyxl
import pytest

from asa_mapper.asa_template import TemplateError, check_template, fill_template, final_asas, template_keys
from asa_mapper.export_parser import parse_export
from tests.helpers import make_template_xlsx


def test_template_keys(template_xlsx):
    keys = template_keys(template_xlsx)
    assert "bank" in keys and "bank#Yes" in keys and "investortype#entity" in keys and "sig" in keys


def test_rejects_non_template(library_xlsx):
    with pytest.raises(TemplateError, match="Mapped Asa"):
        template_keys(library_xlsx)
    with pytest.raises(TemplateError):
        template_keys(b"nope")


def test_check_template_against_form(export_dir):
    form = parse_export(export_dir)
    ok = check_template(make_template_xlsx(), form)
    assert ok == {"matched": 15, "missing_in_template": [], "not_in_form": []}
    bad = check_template(make_template_xlsx(drop={"email1", "bank#No"}, extra={"ghost", "bank#Maybe"}), form)
    assert bad["missing_in_template"] == ["bank#No", "email1"]
    assert bad["not_in_form"] == ["bank#Maybe", "ghost"]


def test_check_template_rejects_other_form(export_dir):
    with pytest.raises(TemplateError, match="does not match"):
        check_template(make_template_xlsx(drop=set(template_keys(make_template_xlsx()))), parse_export(export_dir))


def test_fill_template_sets_only_mapped_asa(template_xlsx):
    data, unknown = fill_template(template_xlsx, {"bank": "asa_bhc_specialentities",
                                                  "bank#Yes": "asa_yes_bhc_specialentities", "ghost": "asa_x"})
    assert unknown == ["ghost"]
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    header = [c.value for c in ws[1]]
    rows = {(r[1] or "") + ("#" + r[2] if r[2] else ""): r for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["bank"][header.index("Mapped Asa")] == "asa_bhc_specialentities"
    assert rows["#Yes"][header.index("Mapped Asa")] == "asa_yes_bhc_specialentities"
    assert rows["email1"][header.index("Mapped Asa")] in (None, "")
    assert rows["bank"][header.index("Alias path")] == "questionnaire<part2<bank"


def test_final_asas_prefers_decisions_then_system_export():
    review = {"existing_source": "asa_export", "fields": [
        {"key": "a", "skip_reason": None, "existing_asa": ["asa_old_a"], "options": [
            {"key": "x", "existing_asa": ["asa_old_ax"]}, {"key": "y", "existing_asa": []}]},
        {"key": "b", "skip_reason": None, "existing_asa": ["asa_old_b"], "options": []},
        {"key": "c", "skip_reason": None, "existing_asa": ["asa_old_c"], "options": []},
    ]}
    decisions = {"a": {"status": "changed", "asa": "asa_new_a",
                       "options": {"y": {"status": "accepted", "asa": "asa_new_ay"}}},
                 "b": {"status": "none", "asa": ""}}
    assert final_asas(review, decisions) == {"a": "asa_new_a", "a#x": "asa_old_ax", "a#y": "asa_new_ay",
                                             "c": "asa_old_c"}
    review["existing_source"] = "pdf"
    assert final_asas(review, decisions) == {"a": "asa_new_a", "a#y": "asa_new_ay"}


def test_fill_template_changes_nothing_but_mapped_asa_cells(template_xlsx):
    import re
    import zipfile

    def parts(data):
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            return {n: zf.read(n) for n in zf.namelist()}

    before = parts(template_xlsx)
    after = parts(fill_template(template_xlsx, {"bank": "asa_bhc_specialentities & co"})[0])
    assert before.keys() == after.keys()
    sheet = next(n for n in before if n.startswith("xl/worksheets/"))
    assert all(before[n] == after[n] for n in before if n != sheet)
    strip_j = lambda x: re.sub(rb'<c r="J\d+"[^>]*?(/>|>.*?</c>)', b"", x)
    assert strip_j(before[sheet]) == strip_j(after[sheet])
    assert b"asa_bhc_specialentities &amp; co" in after[sheet]
